"""Tests for the BEA Use+Make XLSX parser."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import polars as pl
import pytest
from openpyxl.workbook import Workbook


# ---- fixture helper ----


USE_MEMBER = "IOUse_After_Redefinitions_PRO_1997-2023_Summary.xlsx"
MAKE_MEMBER = "IOMake_After_Redefinitions_PRO_1997-2023_Summary.xlsx"


def _populate_sheet(
    ws: Any,
    *,
    title: str,
    year: int,
    col_codes: list[str],
    row_codes: list[str],
    cell_value: float,
    blank_cell: tuple[str, str] | None = None,
) -> None:
    """Write a BEA-style sheet (rows 1-6 boilerplate; rows 7+ data)."""
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value="(Millions of dollars)")
    ws.cell(row=3, column=1, value="Bureau of Economic Analysis")
    ws.cell(row=4, column=1, value=str(year))
    # Row 5: column-header codes in cols 3+
    ws.cell(row=5, column=2, value="Header")
    for j, code in enumerate(col_codes):
        ws.cell(row=5, column=3 + j, value=code)
    # Row 6: column-header names
    ws.cell(row=6, column=1, value="IOCode")
    ws.cell(row=6, column=2, value="Name")
    for j, code in enumerate(col_codes):
        ws.cell(row=6, column=3 + j, value=f"{code} name")
    # Rows 7+: data
    for i, row_code in enumerate(row_codes):
        r = 7 + i
        ws.cell(row=r, column=1, value=row_code)
        ws.cell(row=r, column=2, value=f"{row_code} name")
        for j, col_code in enumerate(col_codes):
            if blank_cell == (row_code, col_code):
                continue  # leave cell empty
            ws.cell(row=r, column=3 + j, value=cell_value)


def _build_synthetic_bea_zip(
    out_path: Path,
    *,
    years: list[int] | None = None,
    industries: list[str] | None = None,
    commodities: list[str] | None = None,
    use_extra_final_demand: list[str] | None = None,
    use_extra_value_added: list[str] | None = None,
    make_extra_invalid_rows: list[str] | None = None,
    use_value: float = 1.0,
    make_value: float = 2.0,
    blank_use_cell: tuple[str, str] | None = None,
) -> Path:
    """Build a synthetic AllTablesIO-shaped zip at `out_path`.

    Use sheet rows = commodities + use_extra_value_added (so VA rows can
    test the row-filter); columns = industries + use_extra_final_demand
    (so FD columns can test the column-filter).

    Make sheet rows = industries + make_extra_invalid_rows (the latter
    used to test pattern-filtering of footnote-text rows that BEA
    sometimes includes at the bottom of Make sheets); columns =
    commodities (Make's columns are always clean).
    """
    years = years or [2022, 2023]
    industries = industries or ["I1", "I2", "I3"]
    commodities = commodities or ["C1", "C2", "C3", "C4"]
    use_extra_final_demand = use_extra_final_demand or []
    use_extra_value_added = use_extra_value_added or []
    make_extra_invalid_rows = make_extra_invalid_rows or []

    use_industry_codes = list(industries) + list(use_extra_final_demand)
    use_commodity_codes = list(commodities) + list(use_extra_value_added)
    make_row_codes = list(industries) + list(make_extra_invalid_rows)

    use_wb = Workbook()
    use_wb.remove(use_wb.active)
    make_wb = Workbook()
    make_wb.remove(make_wb.active)

    for year in years:
        _populate_sheet(
            use_wb.create_sheet(str(year)),
            title="The Use of Commodities by Industries",
            year=year,
            col_codes=use_industry_codes,
            row_codes=use_commodity_codes,
            cell_value=use_value,
            blank_cell=blank_use_cell,
        )
        _populate_sheet(
            make_wb.create_sheet(str(year)),
            title="The Make of Commodities by Industries",
            year=year,
            col_codes=commodities,
            row_codes=make_row_codes,
            cell_value=make_value,
        )

    use_buf = io.BytesIO()
    use_wb.save(use_buf)
    make_buf = io.BytesIO()
    make_wb.save(make_buf)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out_path, "w") as zf:
        zf.writestr(USE_MEMBER, use_buf.getvalue())
        zf.writestr(MAKE_MEMBER, make_buf.getvalue())
    return out_path


# ---- tests ----


def test_industries_and_commodities_from_make(tmp_path: Path) -> None:
    """Make's row codes are the industries; its column codes are the commodities."""
    from eia.sources.bea_io import _extract_make_to_temp, _industries_and_commodities_from_make

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2022, 2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
    )
    make_path = _extract_make_to_temp(zip_path, tmp_path / "_make.xlsx")
    industries, commodities = _industries_and_commodities_from_make(make_path)

    assert industries == ["I1", "I2", "I3"]
    assert commodities == ["C1", "C2", "C3", "C4"]


def test_parse_make_year(tmp_path: Path) -> None:
    """Make is rows=industries × cols=commodities; output is long-form."""
    from datetime import datetime, timezone

    from eia.sources.bea_io import _extract_make_to_temp, _parse_make_year

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        make_value=7.0,
    )
    make_path = _extract_make_to_temp(zip_path, tmp_path / "_make.xlsx")
    fetched_at = datetime(2026, 5, 9, tzinfo=timezone.utc)

    df = _parse_make_year(make_path, year=2023, fetched_at=fetched_at)

    # 3 industries × 4 commodities = 12 rows.
    assert df.height == 12
    assert df.columns == [
        "table_year",
        "industry_code",
        "commodity_code",
        "value_millions",
        "fetched_at",
    ]
    assert df["table_year"].unique().to_list() == [2023]
    assert sorted(df["industry_code"].unique().to_list()) == ["I1", "I2", "I3"]
    assert sorted(df["commodity_code"].unique().to_list()) == ["C1", "C2", "C3", "C4"]
    # All values are the synthetic 7.0 we passed.
    assert df["value_millions"].to_list() == [7.0] * 12


def test_parse_use_year_filters_final_demand_columns(tmp_path: Path) -> None:
    """A column code not in the industries set (final-demand) is dropped from output."""
    from datetime import datetime, timezone

    from eia.sources.bea_io import _extract_use_to_temp, _parse_use_year

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        use_extra_final_demand=["F010"],   # PCE — should be filtered out
    )
    use_path = _extract_use_to_temp(zip_path, tmp_path / "_use.xlsx")
    fetched_at = datetime(2026, 5, 9, tzinfo=timezone.utc)

    df = _parse_use_year(
        use_path,
        year=2023,
        industries={"I1", "I2", "I3"},
        commodities={"C1", "C2", "C3", "C4"},
        fetched_at=fetched_at,
    )

    # 4 commodities × 3 industries = 12 rows; F010 column is dropped.
    assert df.height == 12
    assert "F010" not in df["industry_code"].unique().to_list()


def test_parse_use_year_filters_value_added_rows(tmp_path: Path) -> None:
    """A row code not in the commodities set (value-added) is dropped from output."""
    from datetime import datetime, timezone

    from eia.sources.bea_io import _extract_use_to_temp, _parse_use_year

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        use_extra_value_added=["V001"],   # compensation of employees — filtered out
    )
    use_path = _extract_use_to_temp(zip_path, tmp_path / "_use.xlsx")
    fetched_at = datetime(2026, 5, 9, tzinfo=timezone.utc)

    df = _parse_use_year(
        use_path,
        year=2023,
        industries={"I1", "I2", "I3"},
        commodities={"C1", "C2", "C3", "C4"},
        fetched_at=fetched_at,
    )

    # 4 commodities × 3 industries = 12 rows; V001 row is dropped.
    assert df.height == 12
    assert "V001" not in df["commodity_code"].unique().to_list()


def test_parse_use_year_blank_cell_becomes_null(tmp_path: Path) -> None:
    """A blank Use cell maps to a null value_millions in the long-form output."""
    from datetime import datetime, timezone

    from eia.sources.bea_io import _extract_use_to_temp, _parse_use_year

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        blank_use_cell=("C2", "I2"),
    )
    use_path = _extract_use_to_temp(zip_path, tmp_path / "_use.xlsx")
    fetched_at = datetime(2026, 5, 9, tzinfo=timezone.utc)

    df = _parse_use_year(
        use_path,
        year=2023,
        industries={"I1", "I2", "I3"},
        commodities={"C1", "C2", "C3", "C4"},
        fetched_at=fetched_at,
    )

    target = df.filter(
        (pl.col("commodity_code") == "C2") & (pl.col("industry_code") == "I2")
    )
    assert target.height == 1
    assert target["value_millions"][0] is None


def test_to_cleaned_writes_use_and_make_parquets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """to_cleaned writes both use_summary.parquet and make_summary.parquet, multi-year stacked."""
    from eia.sources.bea_io import BEAIO

    monkeypatch.setattr("eia.config.settings.eia_data_root", tmp_path)
    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2022, 2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        use_extra_final_demand=["F010"],   # Use also has FD col
        use_extra_value_added=["V001"],    # Use also has VA row
    )

    src = BEAIO(year=2023)
    use_path = src.to_cleaned(zip_path)

    assert use_path.exists()
    assert use_path.name == "use_summary.parquet"
    make_path = use_path.parent / "make_summary.parquet"
    assert make_path.exists()

    use_df = pl.read_parquet(use_path)
    # 2 years × 4 commodities × 3 industries = 24 rows (FD col + VA row dropped).
    assert use_df.height == 24
    assert sorted(use_df["table_year"].unique().to_list()) == [2022, 2023]
    assert sorted(use_df["industry_code"].unique().to_list()) == ["I1", "I2", "I3"]
    assert sorted(use_df["commodity_code"].unique().to_list()) == ["C1", "C2", "C3", "C4"]

    make_df = pl.read_parquet(make_path)
    # 2 years × 3 industries × 4 commodities = 24 rows.
    assert make_df.height == 24


def test_load_drops_manifest_and_populates_use_and_make(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """load() registers bea_io_use and bea_io_make and drops bea_io_manifest."""
    from eia.sources.bea_io import BEAIO
    from eia.warehouse.duckdb_impl import DuckDBWarehouse

    monkeypatch.setattr("eia.config.settings.eia_data_root", tmp_path)
    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
    )

    # Build a tmp warehouse that already has the bea_io_manifest stub table to
    # prove load() drops it.
    wh = DuckDBWarehouse(path=tmp_path / "test.duckdb")
    wh.execute_sql(
        "CREATE TABLE bea_io_manifest (filename VARCHAR, table_year SMALLINT, fetched_at TIMESTAMP)"
    )
    wh.execute_sql(
        "CREATE TABLE bea_io_use (table_year SMALLINT, industry_code VARCHAR, commodity_code VARCHAR, "
        "value_millions DOUBLE, fetched_at TIMESTAMP)"
    )
    wh.execute_sql(
        "CREATE TABLE bea_io_make (table_year SMALLINT, industry_code VARCHAR, commodity_code VARCHAR, "
        "value_millions DOUBLE, fetched_at TIMESTAMP)"
    )
    assert "bea_io_manifest" in wh.list_tables()

    src = BEAIO(year=2023)
    use_path = src.to_cleaned(zip_path)
    src.load(use_path, wh)

    tables = wh.list_tables()
    assert "bea_io_manifest" not in tables
    assert "bea_io_use" in tables
    assert "bea_io_make" in tables

    n_use = wh.query("SELECT COUNT(*) AS n FROM bea_io_use")["n"][0]
    n_make = wh.query("SELECT COUNT(*) AS n FROM bea_io_make")["n"][0]
    assert n_use == 12  # 1 year × 4 commodities × 3 industries
    assert n_make == 12  # 1 year × 3 industries × 4 commodities


def test_make_drops_non_iocode_rows(tmp_path: Path) -> None:
    """Footnote-text rows in Make (non-IOCode patterns) are filtered out.

    BEA's published Make sheets sometimes include trailing footnote text
    in the same row layout as data rows (e.g. "Note. Detail may not add to
    total due to rounding."). Without filtering, these get treated as
    industries — observed in real 2023 BEA data.
    """
    from datetime import datetime, timezone

    from eia.sources.bea_io import (
        _extract_make_to_temp,
        _industries_and_commodities_from_make,
        _parse_make_year,
    )

    zip_path = _build_synthetic_bea_zip(
        tmp_path / "AllTablesIO.zip",
        years=[2023],
        industries=["I1", "I2", "I3"],
        commodities=["C1", "C2", "C3", "C4"],
        make_extra_invalid_rows=[
            "Note. Detail may not add to total due to rounding.",
            "1. Consists of only scrap in the use table.",
        ],
    )
    make_path = _extract_make_to_temp(zip_path, tmp_path / "_make.xlsx")

    industries, commodities = _industries_and_commodities_from_make(make_path)
    # Footnote-text rows must NOT appear in the canonical industries set.
    assert industries == ["I1", "I2", "I3"]

    df = _parse_make_year(
        make_path, year=2023, fetched_at=datetime(2026, 5, 9, tzinfo=timezone.utc)
    )
    # 3 industries × 4 commodities = 12 rows; the 2 footnote rows are dropped.
    assert df.height == 12
    assert sorted(df["industry_code"].unique().to_list()) == ["I1", "I2", "I3"]
